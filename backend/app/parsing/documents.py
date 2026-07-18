"""PDF / DOCX text extraction with page or paragraph locators."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pdfplumber
from docx import Document


@dataclass(frozen=True)
class DocumentChunk:
    source_path: str  # path relative to dossier root
    locator: str  # e.g. page:2 or para:12
    text: str


def extract_pdf(path: Path, relative: str) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                chunks.append(
                    DocumentChunk(source_path=relative, locator=f"page:{i}", text=text)
                )
    return chunks


def extract_docx(path: Path, relative: str) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    doc = Document(str(path))
    for i, para in enumerate(doc.paragraphs, start=1):
        text = (para.text or "").strip()
        if text:
            chunks.append(
                DocumentChunk(source_path=relative, locator=f"para:{i}", text=text)
            )
    # Tables as additional locators
    for t_i, table in enumerate(doc.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        body = "\n".join(rows).strip()
        if body:
            chunks.append(
                DocumentChunk(
                    source_path=relative, locator=f"table:{t_i}", text=body
                )
            )
    return chunks


def extract_documents(root: Path) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        rel = str(path.relative_to(root)).replace("\\", "/")
        if suffix == ".pdf":
            chunks.extend(extract_pdf(path, rel))
        elif suffix == ".docx":
            chunks.extend(extract_docx(path, rel))
    return chunks
